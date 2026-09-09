import openpyxl, pypdf, sys
base = r"G:\我的雲端硬碟\catalogue"
# 1) Mitsubishi capacity info.xlsx
wb = openpyxl.load_workbook(base + r"\Mitsubishi\capacity info.xlsx", data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"--- sheet '{ws.title}' dims {ws.dimensions}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [str(c) if c is not None else "" for c in row]
        if any(vals):
            print(" | ".join(vals))
        if i > 18: break
